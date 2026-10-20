"""
Excel (.xlsx) and Word (.docx) report generation for student observation
history, using openpyxl and python-docx.
"""

from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from docx import Document

HEADERS = [
    "차시(세션 ID)",
    "턴",
    "시간(UTC)",
    "이미지",
    "비교모드",
    "학생 질문",
    "AI 답변",
    "객관성 점수",
    "관찰 다양성",
    "과학 용어 수",
    "공간 범위",
    "적용 규칙 ID",
    "AI 단계",
    "모드",
    "입력 종류",
    "지원 수준",
    "호출한 도구",
    "도구 누락",
]

MODE_LABELS = {"assess": "관찰력 진단", "offload": "오프로딩 진단", "support": "AI Agent 지원"}


def sanitize_sheet_name(name: str) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", name or "sheet")
    cleaned = cleaned.strip() or "sheet"
    return cleaned[:31]


def _row_values(row: dict[str, Any]) -> list[Any]:
    images = row.get("images") or []
    return [
        row.get("session_id"),
        row.get("turn_index"),
        row.get("timestamp"),
        ", ".join(images) if isinstance(images, list) else images,
        "예" if row.get("comparison_mode") else "아니오",
        row.get("question"),
        row.get("answer"),
        row.get("objectivity_score"),
        row.get("variety_count"),
        row.get("term_count"),
        row.get("spatial_scope"),
        row.get("pedagogy_rule_id"),
        row.get("ai_stage"),
        MODE_LABELS.get(row.get("mode"), row.get("mode")),
        row.get("input_kind"),
        row.get("support_level"),
        row.get("tools_called"),
        ("예" if row.get("tool_skipped") else "") if row.get("tool_skipped") is not None else "",
    ]


def _autosize(ws) -> None:
    widths: dict[str, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = len(str(cell.value))
            col = cell.column_letter
            widths[col] = max(widths.get(col, 8), min(length + 2, 60))
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _add_sheet(wb: Workbook, title: str, rows: list[dict[str, Any]]):
    ws = wb.create_sheet(title=sanitize_sheet_name(title))
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(_row_values(row))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(ws)
    return ws


def build_student_xlsx(student_id: str, rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    _add_sheet(wb, student_id, rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


TOOL_CALL_HEADERS = [
    "학번",
    "tool 이름",
    "호출 시각(UTC)",
    "입력(관찰문 등)",
    "매칭된 규칙 ID",
    "AI 단계",
    "반환된 힌트/요약",
]

TOOL_NAME_LABELS = {
    "pedagogy_hint": "교육학 안내 조회",
    "history_lookup": "학생 관찰 이력 조회",
    "observation_assessment": "관찰력 평가 데이터 조회",
    "offloading_diagnosis": "오프로딩 단계 진단 조회",
    "support_methods": "오프로딩 단계별 지원방법 조회",
    "observation_methods": "관찰 유형 및 방법 조회",
    "dialogue_log": "개인별 대화 로그 조회",
    "tool_skipped": "도구 호출 누락(자동 기록)",
}


def _add_tool_calls_sheet(wb: Workbook, tool_calls: list[dict[str, Any]]):
    ws = wb.create_sheet(title="Tool호출이력")
    ws.append(TOOL_CALL_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in tool_calls:
        ws.append(
            [
                row.get("student_id"),
                TOOL_NAME_LABELS.get(row.get("tool_name"), row.get("tool_name")),
                row.get("timestamp"),
                row.get("input_text"),
                row.get("pedagogy_rule_id"),
                row.get("ai_stage"),
                row.get("result_summary"),
            ]
        )
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(ws)
    return ws


DIAGNOSIS_HEADERS = [
    "학번", "모드", "종류", "기록 시각(UTC)", "관찰력 수준", "오프로딩 단계", "수준/단계 이름",
    "점수(OQ_A)", "제출 문장 수", "범위 밖 문장 수", "다양도 D", "정량 문장 수", "사용한 관찰 유형",
    "AI 요청 수", "관찰 대행 비율", "AI 반영 비율", "사진 일치율", "수준 판정 출처", "LLM 판정 상태",
    "LLM 판정 근거", "도구 누락 턴 수", "지원 수준",
]


def _add_diagnoses_sheet(wb: Workbook, diagnoses: list[dict[str, Any]]):
    ws = wb.create_sheet(title="진단결과")
    ws.append(DIAGNOSIS_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for d in diagnoses:
        p = d.get("payload") or {}
        judge = p.get("judge") or {}
        ws.append(
            [
                d.get("student_id"),
                MODE_LABELS.get(d.get("mode"), d.get("mode")),
                d.get("kind"),
                d.get("created_at"),
                p.get("level"),
                p.get("stage"),
                p.get("levelLabel") or p.get("stageLabel"),
                p.get("score"),
                p.get("sentenceCount") if p.get("sentenceCount") is not None else p.get("submittedSentenceCount"),
                p.get("outOfScopeMessageCount"),
                p.get("varietyCount"),
                p.get("quantitySentenceCount"),
                ", ".join(str(x) for x in (p.get("typesUsed") or [])),
                p.get("requestCount"),
                p.get("delegateShare"),
                p.get("reuseRatio"),
                judge.get("photoMatchRate"),
                p.get("finalLevelSource") or p.get("finalStageSource") or p.get("method"),
                p.get("judgeStatus"),
                judge.get("rationale"),
                p.get("toolSkippedTurns"),
                ", ".join(str(x) for x in (p.get("supportLevels") or [])),
            ]
        )
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(ws)
    return ws


OUT_OF_SCOPE_LABELS = {"nonVisualSense": "시각 이외 감각", "timeFlow": "시간 흐름"}


def _unit_text(unit: dict[str, Any]) -> str:
    if unit.get("outOfScope"):
        return "범위 밖(" + OUT_OF_SCOPE_LABELS.get(unit["outOfScope"], unit["outOfScope"]) + ")"
    if not unit.get("do"):
        return "0"
    meas = {"qualitative": "정성", "quantitative": "정량"}.get(unit.get("measurement"), "")
    comp = {"single": "단수", "multiple": "다수"}.get(unit.get("comparison"), "")
    scope = {"whole": "전체", "part": "부분"}.get(unit.get("scope"), "")
    return f"{unit['do']}-{meas}-{comp}-{scope}"


def _add_classification_sheet(wb: Workbook, diagnoses: list[dict[str, Any]]):
    """탭 1(관찰력 진단)과 탭 3(지원 후 관찰)의 문장별 분류 결과를 한 줄씩 적는다."""
    ws = wb.create_sheet(title="관찰분류")
    ws.append(
        ["학번", "모드", "번호", "관찰 문장", "사진", "LE", "정보 단위(객관도-측정-비교-범위)",
         "ΣDO", "LE×ΣDO", "범위 밖 표시", "분류 출처"]
    )
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for d in diagnoses:
        p = d.get("payload") or {}
        for i, m in enumerate(p.get("messages") or [], start=1):
            if not isinstance(m, dict) or "units" not in m:
                continue
            ws.append(
                [
                    d.get("student_id"),
                    MODE_LABELS.get(d.get("mode"), d.get("mode")),
                    i,
                    m.get("text"),
                    m.get("image"),
                    m.get("le"),
                    " / ".join(_unit_text(u) for u in m.get("units") or []),
                    m.get("sumDO"),
                    m.get("contribution"),
                    ", ".join(OUT_OF_SCOPE_LABELS.get(f, f) for f in m.get("outOfScope") or []),
                    {"agent": "Agent", "pending": "분류 전"}.get(m.get("source"), "미분류"),
                ]
            )
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(ws)
    return ws


def _add_mode_sessions_sheet(wb: Workbook, mode_sessions: list[dict[str, Any]]):
    ws = wb.create_sheet(title="모드세션")
    ws.append(["학번", "모드", "시작(UTC)", "마감(UTC)", "제출(UTC)", "자동 제출", "사진", "제출 문장 수"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for m in mode_sessions:
        ws.append(
            [
                m.get("student_id"),
                MODE_LABELS.get(m.get("mode"), m.get("mode")),
                m.get("started_at"),
                m.get("deadline_at"),
                m.get("submitted_at"),
                "예" if m.get("auto_submitted") else "아니오",
                ", ".join(m.get("images") or []),
                len(m.get("sentences") or []),
            ]
        )
    _autosize(ws)
    return ws


def _add_sentences_sheet(wb: Workbook, mode_sessions: list[dict[str, Any]]):
    ws = wb.create_sheet(title="제출문장")
    ws.append(["학번", "모드", "번호", "시작 후 경과(초)", "관찰 문장", "사진"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for m in mode_sessions:
        for i, sent in enumerate(m.get("sentences") or [], start=1):
            ws.append(
                [
                    m.get("student_id"),
                    MODE_LABELS.get(m.get("mode"), m.get("mode")),
                    i,
                    sent.get("t") if isinstance(sent, dict) else None,
                    sent.get("text") if isinstance(sent, dict) else str(sent),
                    sent.get("image") if isinstance(sent, dict) else None,
                ]
            )
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(ws)
    return ws


def build_admin_xlsx(
    summary: list[dict[str, Any]],
    rows_by_student: dict[str, list[dict[str, Any]]],
    tool_calls: list[dict[str, Any]] | None = None,
    diagnoses: list[dict[str, Any]] | None = None,
    mode_sessions: list[dict[str, Any]] | None = None,
) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    if diagnoses:
        _add_diagnoses_sheet(wb, diagnoses)
        _add_classification_sheet(wb, diagnoses)
    if mode_sessions:
        _add_mode_sessions_sheet(wb, mode_sessions)
        _add_sentences_sheet(wb, mode_sessions)
    ws_summary = wb.create_sheet(title="학생목록")
    ws_summary.append(
        [
            "학번",
            "대화 턴 수",
            "참여 세션 수",
            "이력조회 tool 호출 수",
            "교육학조회 tool 호출 수",
            "첫 대화",
            "마지막 대화",
        ]
    )
    for cell in ws_summary[1]:
        cell.font = Font(bold=True)
    for s in summary:
        ws_summary.append(
            [
                s.get("student_id"),
                s.get("turn_count"),
                s.get("session_count"),
                s.get("history_lookup_calls", 0),
                s.get("pedagogy_hint_calls", 0),
                s.get("first_at"),
                s.get("last_at"),
            ]
        )
    _autosize(ws_summary)
    if tool_calls:
        _add_tool_calls_sheet(wb, tool_calls)
    for student_id, rows in rows_by_student.items():
        _add_sheet(wb, student_id, rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _add_student_section(doc: Document, student_id: str, rows: list[dict[str, Any]]) -> None:
    doc.add_heading(f"학생 {student_id} 관찰 대화 기록", level=1)
    if not rows:
        doc.add_paragraph("저장된 대화 기록이 없습니다.")
        return
    for row in rows:
        images = row.get("images") or []
        images_text = ", ".join(images) if isinstance(images, list) else (images or "-")
        doc.add_heading(
            f"[{row.get('session_id')}] 턴 {row.get('turn_index')} · {row.get('timestamp')}",
            level=2,
        )
        doc.add_paragraph(
            f"이미지: {images_text}  |  비교모드: {'예' if row.get('comparison_mode') else '아니오'}"
        )
        doc.add_paragraph(
            "객관성 점수: {0}  |  관찰 다양성: {1}  |  과학 용어 수: {2}  |  공간 범위: {3}  |  적용 규칙: {4} ({5}단계)".format(
                row.get("objectivity_score"),
                row.get("variety_count"),
                row.get("term_count"),
                row.get("spatial_scope"),
                row.get("pedagogy_rule_id") or "-",
                row.get("ai_stage") or "-",
            )
        )
        p = doc.add_paragraph()
        p.add_run("학생 질문: ").bold = True
        p.add_run(row.get("question") or "")
        p2 = doc.add_paragraph()
        p2.add_run("AI 답변: ").bold = True
        p2.add_run(row.get("answer") or "")
        doc.add_paragraph("")


def build_student_docx(student_id: str, rows: list[dict[str, Any]]) -> bytes:
    doc = Document()
    _add_student_section(doc, student_id, rows)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_admin_docx(
    summary: list[dict[str, Any]],
    rows_by_student: dict[str, list[dict[str, Any]]],
    tool_calls: list[dict[str, Any]] | None = None,
) -> bytes:
    doc = Document()
    doc.add_heading("전체 학생 관찰 대화 기록", level=0)
    for s in summary:
        doc.add_paragraph(
            "학번 {0}: 총 {1}턴, {2}개 세션 참여, tool 호출 - 이력조회 {3}회 / 교육학조회 {4}회 "
            "(첫 대화 {5} ~ 마지막 대화 {6})".format(
                s.get("student_id"),
                s.get("turn_count"),
                s.get("session_count"),
                s.get("history_lookup_calls", 0),
                s.get("pedagogy_hint_calls", 0),
                s.get("first_at"),
                s.get("last_at"),
            )
        )
    if tool_calls:
        doc.add_page_break()
        doc.add_heading("Tool 호출 이력 (전체)", level=1)
        for row in tool_calls:
            doc.add_paragraph(
                "[{0}] {1} · 학번 {2}\n입력: {3}\n결과: {4} (규칙 {5}, {6}단계)".format(
                    row.get("timestamp"),
                    TOOL_NAME_LABELS.get(row.get("tool_name"), row.get("tool_name")),
                    row.get("student_id"),
                    row.get("input_text") or "-",
                    row.get("result_summary") or "-",
                    row.get("pedagogy_rule_id") or "-",
                    row.get("ai_stage") or "-",
                )
            )
    doc.add_page_break()
    student_ids = list(rows_by_student.keys())
    for idx, student_id in enumerate(student_ids):
        _add_student_section(doc, student_id, rows_by_student[student_id])
        if idx < len(student_ids) - 1:
            doc.add_page_break()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()